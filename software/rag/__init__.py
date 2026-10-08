"""Selective RAG for Naranjo Q2, Q5 and Q10.

Contents: drug-class resolution (openFDA -> RxNorm -> DeepSeek, cached), expert knowledge
base loading and TF-IDF retrieval indices, prompt text, and the SelectiveRAG class.
"""
import os
import re as _re
import json
import time
from collections import defaultdict

import numpy as np
import openpyxl
import requests as _req
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from .. import config
from ..config import call_deepseek, call_deepseek_json


# ════════════════════════════════════════════════════════════════════════════
# Drug-class resolution
# ════════════════════════════════════════════════════════════════════════════


def normalise_class(raw: str) -> str:
    s = _re.sub(r"\s*\[EPC\]", "", raw).strip()
    s = _re.sub(r"[^\w\s]", "", s)
    s = _re.sub(r"\s+", "_", s)
    return s.upper() or "UNKNOWN"


def openfda_epc(drug_name: str):
    primary = drug_name.split("\\")[0].strip().lower()
    names   = ([drug_name.lower()] if drug_name.lower() != primary else []) + [primary]
    for name in names:
        for field in ("openfda.brand_name", "openfda.substance_name",
                      "openfda.generic_name"):
            try:
                r = _req.get("https://api.fda.gov/drug/label.json",
                             params={"search": f'{field}:"{name}"', "limit": 1},
                             timeout=10)
                if r.status_code != 200:
                    continue
                epc = r.json().get("results", [{}])[0].get("openfda", {}) \
                               .get("pharm_class_epc", [])
                if epc:
                    return normalise_class(epc[0])
            except Exception:
                pass
    return None




def rxnorm_class(drug_name: str):
    primary = drug_name.split("\\")[0].strip().lower()
    try:
        r = _req.get("https://rxnav.nlm.nih.gov/REST/rxcui.json",
                     params={"name": primary, "search": 2}, timeout=10)
        if r.status_code != 200:
            return None
        cuis = r.json().get("idGroup", {}).get("rxnormId", [])
        if not cuis:
            return None
        for source in ("ATC", "VA", "MESH"):
            r2 = _req.get(
                "https://rxnav.nlm.nih.gov/REST/rxclass/class/byRxcui.json",
                params={"rxcui": cuis[0], "relaSource": source}, timeout=10)
            items = (r2.json().get("rxclassDrugInfoList", {})
                              .get("rxclassDrugInfo", [])
                     if r2.status_code == 200 else [])
            if items:
                best = max(items,
                           key=lambda x: len(x["rxclassMinConceptItem"]["classId"]))
                return normalise_class(best["rxclassMinConceptItem"]["className"])
    except Exception:
        pass
    return None

def deepseek_batch(drugs: list) -> dict:
    if not drugs:
        return {}
    prompt = (
        "You are a senior pharmacologist. For each drug below return its "
        "pharmacological class as a concise uppercase label with underscores "
        "(e.g. SULFONYLUREA, BRAF_INHIBITOR, ANTI_TNF_BIOLOGIC, INSULIN, SSRI).\n"
        "For combination products use the primary active component class.\n"
        "Return ONLY a JSON object mapping each exact drug name to its class.\n\n"
        "Drugs:\n" + "\n".join(f"- {d}" for d in drugs)
    )
    try:
        parsed = call_deepseek_json(prompt, system="Return only valid JSON, no other text.", timeout=90)
        return {k: normalise_class(str(v)) for k, v in parsed.items() if k in drugs}
    except Exception as e:
        print(f"    ⚠️  DeepSeek batch error: {e}")
        return {}



BATCH_SIZE = 30
FDA_PAUSE = 0.25
RXNORM_PAUSE = 0.12


def build_class_map(unique_drugs: list, cache_path=None) -> dict:
    """
    Resolve drug → class for all unique_drugs.
    Loads from cache if available; only queries APIs for new drugs.
    """
    cache_path = str(cache_path or config.DRUG_CLASS_CACHE_PATH)
    cache: dict = {}
    if os.path.exists(cache_path):
        with open(cache_path) as f:
            cache = json.load(f)
        to_resolve = [d for d in unique_drugs if d.upper() not in cache]
        if not to_resolve:
            print(f"[DrugClass] ✓ All {len(unique_drugs)} drugs in cache.")
            return cache
        print(f"[DrugClass] Cache hit for "
              f"{len(unique_drugs)-len(to_resolve)}/{len(unique_drugs)} drugs. "
              f"Resolving {len(to_resolve)} new ones…")
    else:
        to_resolve = unique_drugs
        print(f"[DrugClass] No cache. Resolving {len(to_resolve)} drugs…")

    unresolved = []
    hits1 = 0
    print("  [1/3] OpenFDA…")
    for i, drug in enumerate(to_resolve, 1):
        cls = openfda_epc(drug)
        if cls:
            cache[drug.upper()] = cls
            hits1 += 1
        else:
            unresolved.append(drug)
        if i % 20 == 0 or i == len(to_resolve):
            print(f"    {i}/{len(to_resolve)}  ({hits1} resolved so far)")
        time.sleep(FDA_PAUSE)
    print(f"  OpenFDA: {hits1} resolved, {len(unresolved)} remain.")

    still = []
    hits2 = 0
    print("  [2/3] RxNorm…")
    for drug in unresolved:
        cls = rxnorm_class(drug)
        if cls:
            cache[drug.upper()] = cls
            hits2 += 1
        else:
            still.append(drug)
        time.sleep(RXNORM_PAUSE)
    print(f"  RxNorm: {hits2} resolved, {len(still)} remain.")

    if still:
        hits3 = 0
        print(f"  [3/3] DeepSeek batch ({len(still)} drugs)…")
        for start in range(0, len(still), BATCH_SIZE):
            batch = still[start: start + BATCH_SIZE]
            result = deepseek_batch(batch)
            for drug, cls in result.items():
                cache[drug.upper()] = cls
                hits3 += 1
            for drug in batch:
                if drug.upper() not in cache:
                    cache[drug.upper()] = "UNKNOWN"
        print(f"  DeepSeek: {hits3} resolved.")

    with open(cache_path, "w") as f:
        json.dump(cache, f, indent=2, sort_keys=True)
    print(f"[DrugClass] Cache saved → {cache_path}")
    return cache


# ════════════════════════════════════════════════════════════════════════════
# Expert knowledge base and retrieval indices
# ════════════════════════════════════════════════════════════════════════════


# ── KB column positions (0-based) ────────────────────────────────────────────
_COL_CASE_ID = 1
_COL_DRUG    = 3
_COL_PT      = 2
_COL_NARR    = 47
_COL_META    = 51
_RAG_Q = {
    "Q2":  {"h_score": 10, "h_reason": 11},
    "Q5":  {"h_score": 22, "h_reason": 23},
    "Q10": {"h_score": 42, "h_reason": 43},
}

# ── Missing-date / missing-confounder detectors ──────────────────────────────
_DATE_FIELDS = ["Therapy Start Date", "Therapy End Date", "Event Date",
                "therapy start", "event date"]
_NA_TOKENS   = {"na", "nan", "missing", "unknown", "unk", "n/a", ""}

def has_any_date(ctx: str) -> bool:
    for field in _DATE_FIELDS:
        m = _re.search(rf"{_re.escape(field)}\s*[:\-]\s*([^\n;│]+)",
                       ctx, _re.IGNORECASE)
        if m:
            val = m.group(1).strip().lower()
            if val not in _NA_TOKENS and len(val) >= 4:
                return True
    return False

def has_any_confounder_info(ctx: str) -> bool:
    signals = ["concurrent", "concomitant", "co-medication", "co-morbid",
               "indication for use", "underlying condition", "medical history",
               "confound"]
    return any(s in ctx.lower() for s in signals)


def _named_layout(header_row) -> dict | None:
    """If the first row holds column names (MARCA_vs_expert_comparison layout), map them."""
    names = [str(v).strip() if v is not None else "" for v in header_row]
    if "q2_score_human" not in names:
        return None
    idx = {n: i for i, n in enumerate(names)}
    rag_q = {q: {"h_score": idx[f"{q.lower()}_score_human"], "h_reason": idx[f"{q.lower()}_reasoning_human"]}
             for q in ("Q2", "Q5", "Q10")}
    return {"case_id": idx["case_id"], "pt": idx["pt"], "drug": idx["drug_name"],
            "narr": idx["case_narrative"], "meta": idx["questions2_9_ai_answer"],
            "rag_q": rag_q, "first_row": 2}


def kb_layout(path: str) -> dict:
    """Column positions of the KB workbook.

    Two layouts are supported:
      * named columns in row 1 (e.g. MARCA_vs_expert_comparison.xlsx: q2_score_human, …);
      * the original study workbook without headers in row 1, data from row 3
        (fixed positions defined at the top of this module).
    """
    wb = openpyxl.load_workbook(path, read_only=True)
    header = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
    wb.close()
    return _named_layout(header) or {
        "case_id": _COL_CASE_ID, "pt": _COL_PT, "drug": _COL_DRUG,
        "narr": _COL_NARR, "meta": _COL_META, "rag_q": _RAG_Q, "first_row": 3}


def load_expert_kb(path: str, class_map: dict) -> list:
    if not os.path.exists(path):
        print(f"⚠️  KB not found at {path}")
        return []
    L = kb_layout(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    records = []
    for row in ws.iter_rows(min_row=L["first_row"], values_only=True):
        if not any(row):
            continue
        drug = str(row[L["drug"]] or "").strip()
        pt   = str(row[L["pt"]]   or "").strip()
        if not drug or not pt:
            continue
        narr = str(row[L["narr"]] or "")
        meta = str(row[L["meta"]] or "")
        rec = {
            "case_id":        str(row[L["case_id"]] or "").strip(),
            "drug":           drug,
            "pt":             pt,
            "drug_class":     class_map.get(drug.upper(), "UNKNOWN"),
            "narrative":      narr,
            "metadata":       meta,
            "no_dates":       not has_any_date(narr),
            "no_confounders": not has_any_confounder_info(narr + " " + meta),
        }
        for qname, cols in L["rag_q"].items():
            rec[f"{qname}_h_score"]  = row[cols["h_score"]]
            rec[f"{qname}_h_reason"] = str(row[cols["h_reason"]] or "")
        records.append(rec)
    wb.close()
    print(f"[KB] Loaded {len(records)} records, "
          f"{len({r['drug'] for r in records})} unique drugs.")
    return records


class ExpertVectorDB:
    def __init__(self, records: list):
        self.records = records
        if not records:
            return

        # exact case_id lookup
        self._by_id = {r["case_id"]: r for r in records}

        # ── flat indices (v2 style, global over all records) ──────────────
        ae_corpus  = [r["pt"].lower() for r in records]
        ctx_corpus = [
            f"{r['drug'].lower()} {r['pt'].lower()} {r['narrative'][:400].lower()}"
            for r in records
        ]
        self._ae_vec  = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
        self._ae_mat  = self._ae_vec.fit_transform(ae_corpus)
        self._ctx_vec = TfidfVectorizer(ngram_range=(1, 2), min_df=1, max_features=5000)
        self._ctx_mat = self._ctx_vec.fit_transform(ctx_corpus)

        # ── per-drug flat index (Tier 1) ──────────────────────────────────
        # corpus: pt text only; searched by pt query
        drug_buckets  = defaultdict(list)
        class_buckets = defaultdict(list)
        for r in records:
            drug_buckets[r["drug"].upper()].append(r["pt"].lower())
            if r["drug_class"] != "UNKNOWN":
                class_buckets[r["drug_class"]].append(r["pt"].lower())

        self._drug_recs  = defaultdict(list)
        self._class_recs = defaultdict(list)
        for r in records:
            self._drug_recs[r["drug"].upper()].append(r)
            if r["drug_class"] != "UNKNOWN":
                self._class_recs[r["drug_class"]].append(r)

        self._drug_idx  = {}   # drug_upper  → (vec, mat)
        self._class_idx = {}   # class_upper → (vec, mat)

        for drug_key, pts in drug_buckets.items():
            vec = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
            self._drug_idx[drug_key] = (vec, vec.fit_transform(pts))

        for cls_key, pts in class_buckets.items():
            vec = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
            self._class_idx[cls_key] = (vec, vec.fit_transform(pts))

        # ── missing-date sub-pool (v2 style) ─────────────────────────────
        nd = [r for r in records if r["no_dates"]]
        self._no_date_records = nd
        if nd:
            nd_corpus = [
                f"{r['drug'].lower()} {r['pt'].lower()}" for r in nd
            ]
            self._nd_vec = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
            self._nd_mat = self._nd_vec.fit_transform(nd_corpus)
        else:
            self._nd_vec = self._nd_mat = None

        n_d = len(self._drug_idx)
        n_c = len(self._class_idx)
        print(f"[VectorDB] Built — {n_d} drug buckets, {n_c} class buckets, "
              f"{len(nd)} no-date records.")

    # ── shared scorer (same logic as v2 _query) ───────────────────────────
    def _query(self, mat, vec, pool: list, text: str,
               top_k: int, exclude_id: str) -> list:
        q    = vec.transform([text.lower()])
        sims = cosine_similarity(q, mat).flatten()
        order = np.argsort(sims)[::-1]
        out   = []
        for i in order:
            r = pool[i]
            if exclude_id and r["case_id"] == exclude_id:
                continue
            if sims[i] < 0.01:
                break
            out.append({**r, "_sim": float(sims[i])})
            if len(out) >= top_k:
                break
        return out

    # ── exact lookup ──────────────────────────────────────────────────────
    def lookup_exact(self, case_id) -> dict | None:
        return self._by_id.get(str(case_id).strip())

    # ── two-tier retrieval ────────────────────────────────────────────────
    def query_two_tier(self, drug: str, pt: str,
                       top_k: int = 3, exclude_id: str = "") -> list:
        """
        Tier 1: records with the SAME drug, ranked by AE similarity.
        Tier 2: records with the SAME drug class, ranked by AE similarity.
        Falls back to global context search if both tiers are empty.
        """
        drug_key  = drug.strip().upper()
        drug_class = "UNKNOWN"
        # get class from any record that matches this drug
        for r in self.records:
            if r["drug"].upper() == drug_key:
                drug_class = r["drug_class"]
                break

        results  = []
        seen_ids = set()

        # Tier 1 — same drug, search by PT
        if drug_key in self._drug_idx:
            vec, mat = self._drug_idx[drug_key]
            pool = self._drug_recs[drug_key]
            for r in self._query(mat, vec, pool, pt, top_k, exclude_id):
                if r["case_id"] not in seen_ids:
                    results.append({**r, "_sim_tier": "exact_drug"})
                    seen_ids.add(r["case_id"])

        # Tier 2 — same class, search by PT (fill remaining slots)
        if len(results) < top_k and drug_class not in ("UNKNOWN", "", None):
            if drug_class in self._class_idx:
                vec, mat = self._class_idx[drug_class]
                pool = self._class_recs[drug_class]
                for r in self._query(mat, vec, pool, pt, top_k * 3, exclude_id):
                    if r["case_id"] not in seen_ids:
                        results.append({**r, "_sim_tier": "same_class"})
                        seen_ids.add(r["case_id"])
                    if len(results) >= top_k:
                        break

        # Fallback — global context search (v2 behaviour)
        if not results:
            for r in self._query(self._ctx_mat, self._ctx_vec, self.records,
                                 f"{drug} {pt}", top_k, exclude_id):
                if r["case_id"] not in seen_ids:
                    results.append({**r, "_sim_tier": "global"})
                    seen_ids.add(r["case_id"])

        return results[:top_k]

    # ── AE-only search (for Q10) ──────────────────────────────────────────
    def query_ae(self, pt: str, top_k: int = 2, exclude_id: str = "") -> list:
        return [
            {**r, "_sim_tier": "ae_match"}
            for r in self._query(self._ae_mat, self._ae_vec,
                                 self.records, pt, top_k, exclude_id)
        ]

    # ── missing-date pool search ──────────────────────────────────────────
    def query_missing_dates(self, drug: str, pt: str,
                            top_k: int = 3, exclude_id: str = "") -> list:
        if self._nd_mat is None or not self._no_date_records:
            return []
        return [
            {**r, "_sim_tier": "no_date_pool"}
            for r in self._query(self._nd_mat, self._nd_vec,
                                 self._no_date_records,
                                 f"{drug} {pt}", top_k, exclude_id)
        ]


# ════════════════════════════════════════════════════════════════════════════
# Prompts
# ════════════════════════════════════════════════════════════════════════════
_EXPERT_Q2_NO_DATES = (
    "the case does not provide any information on therapy start and end date "
    "or on the event date, thus temporal and biological plausibility cannot be established."
)
_EXPERT_Q5_NO_INFO = (
    "No information is provided about confounding factors or concurrent risk factors "
    "in the report. Confounding cannot be assessed."
)

Q_DEF = {
    "Q2": (
        "Q2 – Temporal Plausibility\n"
        "Score: 2=clear temporal relationship (drug before AE, plausible latency); "
        "1=unclear or partial; 0=no data (dates missing/not assessable); "
        "-1=implausible (AE before drug).\n"
        "If therapy start date AND event date are both missing/NA, "
        "score MUST be 0 and state temporal plausibility cannot be established."
    ),
    "Q5": (
        "Q5 – Confounding Factors\n"
        "Score: 2=no confounders; 1=minor; 0=no confounder info; "
        "-1=significant confounders; -2=alternative cause clearly more likely.\n"
        "If no concurrent medications and no comorbidities info, score MUST be 0."
    ),
    "Q10": (
        "Q10 – Sign vs Symptom\n"
        "Score: 1=clinical sign (objectively measurable); "
        "0=symptom (subjective/patient-reported); -1=not confirmed."
    ),
}


def build_ref_block(refs, qname):
    lines = []
    for i, r in enumerate(refs, 1):
        tier = r.get("_sim_tier", "")
        tier_label = (
            "same drug"  if tier == "exact_drug"  else
            f"same class · {r.get('drug_class','?')}" if tier == "same_class" else
            "no-date pool" if tier == "no_date_pool" else
            "global match"
        )
        lines.append(
            f"[Ref {i} | {tier_label} | sim={r.get('_sim', 0):.0%} | "
            f"case_id={r['case_id']} | {r['drug']} / {r['pt']}]\n"
            f"  Expert {qname} score  : {r.get(f'{qname}_h_score', '?')}\n"
            f"  Expert {qname} reason : {r.get(f'{qname}_h_reason', '') or '(none)'}"
        )
    return "\n\n".join(lines)


# ════════════════════════════════════════════════════════════════════════════
# Selective RAG
# ════════════════════════════════════════════════════════════════════════════




def _call_llm(prompt_user, system_msg=None):
    """DeepSeek call that returns an error string instead of raising (as in the notebook)."""
    try:
        if system_msg:
            return call_deepseek(prompt_user, system=system_msg)
        return call_deepseek(prompt_user)
    except Exception as e:
        return f"[LLM error: {e}]"


def dual_answer(qname, drug, pt, context, db, case_id=""):
    case_block   = f"Drug: {drug}\nAdverse Event: {pt}\nContext:\n{context[:config.CASE_CONTEXT_CHARS]}"
    no_dates     = not has_any_date(context)
    no_conf_info = not has_any_confounder_info(context)
    expert_exact = db.lookup_exact(case_id)

    # Retrieval
    if qname == "Q10":
        similar = db.query_ae(pt, top_k=2, exclude_id=case_id)
    elif qname == "Q2" and no_dates:
        similar = db.query_missing_dates(drug, pt, top_k=3, exclude_id=case_id)
    else:
        similar = db.query_two_tier(drug, pt, top_k=3, exclude_id=case_id)
        if qname == "Q5" and no_conf_info:
            # prefer refs that also had no confounder info
            pref = [r for r in similar if r.get("no_confounders")]
            rest = [r for r in similar if not r.get("no_confounders")]
            similar = (pref + rest)[:3]

    # No-RAG
    ans_no_rag = _call_llm(
        f"Answer the following pharmacovigilance question.\n\n"
        f"Question:\n{Q_DEF[qname]}\n\n"
        f"Case data:\n{case_block}\n\n"
        f"Reply as:\nScore: <number>\nReasoning: <justification>"
    )

    # With-RAG
    rag_parts = []
    if qname == "Q2" and no_dates:
        rag_parts.append(
            "--- MANDATORY EXPERT RULE: MISSING DATES ---\n"
            "Both therapy start date and event date are NA/missing.\n"
            "Expert annotators consistently score Q2 = 0 and state:\n"
            f'  "{_EXPERT_Q2_NO_DATES}"\n'
            "You MUST apply this rule: score Q2 = 0."
        )
    if qname == "Q5" and no_conf_info:
        rag_parts.append(
            "--- MANDATORY EXPERT RULE: NO CONFOUNDER INFO ---\n"
            "No concurrent medications or comorbidities info available.\n"
            "Expert annotators score Q5 = 0 in this situation.\n"
            "You MUST apply this rule: score Q5 = 0."
        )
    ref_block = build_ref_block(similar, qname)
    if ref_block:
        rag_parts.append(
            f"--- Expert Reference Cases (annotated KB) ---\n{ref_block}\n"
            "--- End References ---\n"
            "Use these to calibrate your answer where relevant."
        )

    if rag_parts:
        ans_with_rag = _call_llm(
            f"Answer the following pharmacovigilance question.\n\n"
            f"Question:\n{Q_DEF[qname]}\n\n"
            f"Case data:\n{case_block}\n\n"
            + "\n\n".join(rag_parts)
            + "\n\nReply as:\nScore: <number>\nReasoning: <justification>"
        )
    else:
        ans_with_rag = ans_no_rag + "\n[No expert context available]"

    return ans_no_rag, ans_with_rag, similar, expert_exact


class SelectiveRAG:
    """Holds the expert KB and its retrieval indices; collects results per case."""

    def __init__(self, records: list):
        self.records = records
        self.db = ExpertVectorDB(records) if records else None
        self.results = []          # one entry per drug-event pair (was `rag_results`)

    @classmethod
    def from_workbook(cls, path=None, cache_path=None):
        """Load the KB workbook, resolve drug classes (cached) and build the indices."""
        path = str(path or config.KB_PATH)
        if not os.path.exists(path):
            print(f"\n❌  KB not found: {path}")
            return cls([])
        layout = kb_layout(path)
        wb = openpyxl.load_workbook(path, read_only=True)
        raw_drugs = [str(row[layout["drug"]]).strip()
                     for row in wb.active.iter_rows(min_row=layout["first_row"], values_only=True)
                     if row[layout["drug"]]]
        wb.close()
        class_map = build_class_map(sorted(set(raw_drugs)), cache_path=cache_path)
        records = load_expert_kb(path, class_map)
        rag = cls(records)
        if records:
            print(f"\n✅ RAG module ready. {len(records)} records | "
                  f"{len({r['drug'] for r in records})} drugs | "
                  f"{len({r.get('drug_class', 'UNKNOWN') for r in records})} classes")
        return rag

    def reset(self):
        """Clear results before assessing a new case."""
        self.results = []

    def run_rag_for_pair(self, drug: str, pt: str, context_text: str, case_id: str = "") -> dict:
        """Called once per drug-AE pair from the main pipeline."""
        if not self.records or self.db is None:
            print("⚠️  run_rag_for_pair: KB not loaded — skipping.")
            return {"drug": drug, "pt": pt, "case_id": case_id}

        entry = {"drug": drug, "pt": pt, "case_id": case_id}
        for qname in ["Q2", "Q5", "Q10"]:
            no_rag, with_rag, refs, exact = dual_answer(qname, drug, pt, context_text, self.db, case_id)
            entry[qname] = {
                "version_no_rag": no_rag,
                "version_with_rag": with_rag,
                "retrieved_refs": refs,
                "expert_exact": {
                    "case_id": exact["case_id"],
                    "h_score": exact.get(f"{qname}_h_score"),
                    "h_reason": exact.get(f"{qname}_h_reason", ""),
                } if exact else None,
            }
        self.results.append(entry)
        return entry
